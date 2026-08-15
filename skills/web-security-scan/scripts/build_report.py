# -*- coding: utf-8 -*-
"""สร้างรายงาน Excel จากผลสแกน (scan.sh) + เช็คลิสต์ 189 ข้อ

usage: python build_report.py scan-example.com-20260815.json --out Report.xlsx
"""
import json, sys, os, argparse, datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

F = "Arial"
NAVY, WHITE = "1F3864", "FFFFFF"
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
RISK_FILL = {"Critical": "FFC7CE", "High": "FFD9A0", "Medium": "FFF2CC", "Low": "E2EFDA"}
RISK_TEXT = {"Critical": "9C0006", "High": "974706", "Medium": "7F6000", "Low": "375623"}
STAT_FILL = {"ไม่ผ่าน": "FFC7CE", "ผ่าน": "C6EFCE", "ยังไม่ตรวจ": "F2F2F2", "ไม่เกี่ยวข้อง": "EDEDED"}
RISK_ORDER = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3}

HERE = os.path.dirname(os.path.abspath(__file__))
CHECKLIST = os.path.join(HERE, "..", "assets", "checklist_189.json")


def evaluate(scan):
    """คืน dict: item_id -> (สถานะ, หลักฐาน)  เฉพาะข้อที่ตรวจอัตโนมัติได้"""
    r = {}
    def set_(i, ok, ev):
        r[i] = ("ผ่าน" if ok else "ไม่ผ่าน", ev)

    def unknown_(i, ev):
        """ตรวจไม่ได้ ≠ ผ่าน.

        A check whose evidence never arrived must not be scored. This used to
        be implicit — an empty protocol list read as "no old protocols found",
        which is a pass on the most dangerous item in the checklist — and the
        owner had no way to tell that apart from a real one.
        """
        r[i] = ("ยังไม่ตรวจ", ev)

    h = {k.lower(): v for k, v in scan.get("headers", {}).items()}
    tls = scan.get("tls", {}) or {}
    paths = scan.get("exposed_paths", {}) or {}
    page = scan.get("page", {}) or {}
    dns = scan.get("dns", {}) or {}
    proj = scan.get("project", {}) or {}
    cookies = scan.get("cookies", []) or []

    # ── TLS ──────────────────────────────────────────
    NO_OPENSSL = "ไม่มี openssl บนเครื่องที่รันสแกน — ตรวจด้วย SSL Labs หรือ nmap --script ssl-enum-ciphers"
    if tls.get("skipped"):
        for i in (1, 3, 6, 11):
            unknown_(i, NO_OPENSSL)
    else:
        protos = tls.get("protocols", [])
        unprobed = tls.get("protocols_unknown", [])
        old = [p for p in protos if p in ("tls1", "tls1_1")]
        old_unprobed = [p for p in unprobed if p in ("tls1", "tls1_1")]
        if old:
            set_(1, False, f"รองรับ: {', '.join(protos)}")
        elif old_unprobed or not protos:
            # The local OpenSSL would not offer these, so the server was never
            # asked. Silence is not a pass.
            unknown_(1, "openssl บนเครื่องนี้ไม่ยอมต่อ " + (", ".join(old_unprobed) or "TLS รุ่นเก่า")
                     + " จึงยังไม่รู้ว่าเซิร์ฟเวอร์เปิดอยู่ไหม — ตรวจด้วย SSL Labs หรือ nmap --script ssl-enum-ciphers")
        else:
            set_(1, True, f"รองรับ: {', '.join(protos)}")

        if tls.get("days_left") not in ("", None):
            d = int(tls["days_left"])
            set_(3, d > 30, f"เหลืออีก {d} วัน (หมดอายุ {tls.get('not_after','')})")
        else:
            unknown_(3, "อ่านวันหมดอายุใบรับรองไม่ได้")

        if tls.get("key_bits"):
            set_(6, int(tls["key_bits"]) >= 2048, f"{tls['key_bits']} bit")
        else:
            unknown_(6, "อ่านขนาดคีย์ไม่ได้")

        # Stapling is only a finding if we got a handshake at all; without one
        # "ไม่มี stapling" describes our own openssl, not the server.
        if tls.get("handshake_ok") is False:
            unknown_(11, "ต่อ TLS ไม่สำเร็จจากเครื่องที่รันสแกน")
        else:
            set_(11, bool(tls.get("ocsp_stapling")), "OCSP stapling " + ("เปิด" if tls.get("ocsp_stapling") else "ปิด"))

    # ── Headers ──────────────────────────────────────
    hsts = h.get("strict-transport-security", "")
    maxage = 0
    if "max-age=" in hsts:
        try: maxage = int(hsts.split("max-age=")[1].split(";")[0])
        except Exception: maxage = 0
    set_(7, maxage >= 31536000, hsts or "ไม่มี HSTS header")
    set_(8, "includesubdomains" in hsts.lower(), hsts or "ไม่มี HSTS header")

    csp = h.get("content-security-policy", "")
    set_(15, bool(csp) and "unsafe-inline" not in csp and "unsafe-eval" not in csp,
         (csp[:150] + "…") if len(csp) > 150 else (csp or "ไม่มี CSP"))
    set_(16, "frame-ancestors" in csp, "มี frame-ancestors" if "frame-ancestors" in csp else "ไม่มี frame-ancestors")
    xfo = h.get("x-frame-options", "")
    set_(17, xfo.upper() in ("DENY", "SAMEORIGIN"), xfo or "ไม่มี X-Frame-Options")
    set_(18, h.get("x-content-type-options", "").lower() == "nosniff",
         h.get("x-content-type-options") or "ไม่มี X-Content-Type-Options")
    set_(19, bool(h.get("referrer-policy")), h.get("referrer-policy") or "ไม่มี Referrer-Policy")
    set_(20, bool(h.get("permissions-policy")), h.get("permissions-policy") or "ไม่มี Permissions-Policy")

    leaks = [f"{k}: {h[k]}" for k in ("server", "x-powered-by", "x-aspnet-version", "x-generator") if h.get(k)]
    verbose = [l for l in leaks if any(ch.isdigit() for ch in l.split(":", 1)[1])]
    set_(21, not verbose, "; ".join(leaks) if leaks else "ไม่พบ header ที่บอกเวอร์ชัน")

    acao = h.get("access-control-allow-origin", "")
    acac = h.get("access-control-allow-credentials", "").lower() == "true"
    set_(22, not (acao == "*" and acac), f"Allow-Origin: {acao or '-'} / Credentials: {acac}")
    set_(27, "charset" in h.get("content-type", "").lower(), h.get("content-type") or "-")

    # ── Redirect / mixed content ─────────────────────
    rd = scan.get("http_redirect", {}) or {}
    set_(9, str(rd.get("location", "")).startswith("https://") and rd.get("status") in (301, 308),
         f"HTTP {rd.get('status')} → {rd.get('location') or '(ไม่ redirect)'}")
    set_(10, page.get("mixed_content_refs", 0) == 0, f"พบ {page.get('mixed_content_refs',0)} รายการที่ยังเป็น http://")

    # ── Cookies ──────────────────────────────────────
    if cookies:
        set_(67, all(c["httponly"] for c in cookies), f"{sum(c['httponly'] for c in cookies)}/{len(cookies)} cookie มี HttpOnly")
        set_(68, all(c["secure"] for c in cookies), f"{sum(c['secure'] for c in cookies)}/{len(cookies)} cookie มี Secure")
        set_(69, all(c["samesite"] for c in cookies), f"{sum(1 for c in cookies if c['samesite'])}/{len(cookies)} cookie มี SameSite")

    # ── Exposed paths ────────────────────────────────
    # A site that answers 200 for a path that cannot exist answers 200 for
    # everything, and "200 with a body" stops meaning "this file is exposed".
    # Every path-derived item below is then unverifiable rather than failed —
    # otherwise the owner is handed three dozen invented Criticals.
    soft404 = bool((scan.get("path_baseline") or {}).get("soft_404"))
    SOFT404_NOTE = ("เว็บนี้ตอบ 200 ให้ path ที่ไม่มีอยู่จริง (soft 404) "
                    "จึงใช้ status 200 ตัดสินไม่ได้ — ต้องเปิดดูเนื้อหาไฟล์เอง")

    def open_(p):
        d = paths.get(p, {})
        return d.get("code") == 200 and d.get("size", 0) > 0

    def path_(i, ok, ev):
        """A verdict that rests on a 200 — only trustworthy when 404s work."""
        if soft404:
            unknown_(i, SOFT404_NOTE)
        else:
            set_(i, ok, ev)

    set_(29, not scan.get("directory_listing_hint"), "พบ Index of" if scan.get("directory_listing_hint") else "ไม่พบ directory listing")
    git = [p for p in ("/.git/config", "/.git/HEAD", "/.svn/entries") if open_(p)]
    path_(30, not git, ", ".join(git) + " เข้าถึงได้" if git else "เข้าถึงไม่ได้")
    envs = [p for p in ("/.env", "/.env.local", "/.env.production", "/config.php.bak", "/.DS_Store", "/wp-config.php.bak") if open_(p)]
    path_(31, not envs, ", ".join(envs) + " เข้าถึงได้" if envs else "เข้าถึงไม่ได้")
    php = [p for p in ("/phpinfo.php", "/info.php", "/test.php") if open_(p)]
    path_(32, not php, ", ".join(php) + " เข้าถึงได้" if php else "เข้าถึงไม่ได้")
    dirs = [p for p in ("/backup/", "/backups/", "/old/", "/temp/", "/tmp/", "/dev/", "/test/") if open_(p)]
    path_(35, not dirs, ", ".join(dirs) + " เข้าถึงได้" if dirs else "เข้าถึงไม่ได้")
    sqls = [p for p in ("/db.sql", "/dump.sql", "/database.sql", "/backup.zip") if open_(p)]
    path_(36, not sqls, ", ".join(sqls) + " เข้าถึงได้" if sqls else "เข้าถึงไม่ได้")
    set_(38, page.get("possible_secret_matches", 0) == 0,
         f"พบข้อความคล้าย key/secret {page.get('possible_secret_matches',0)} จุดใน HTML/JS หน้าแรก")
    path_(40, paths.get("/robots.txt", {}).get("code") == 200, "มี robots.txt" if paths.get("/robots.txt", {}).get("code") == 200 else "ไม่มี robots.txt")
    pkg = [p for p in ("/package.json", "/composer.json", "/composer.lock", "/package-lock.json", "/yarn.lock") if open_(p)]
    path_(41, not pkg, ", ".join(pkg) + " เข้าถึงได้" if pkg else "เข้าถึงไม่ได้")
    path_(128, not open_("/phpmyadmin/") and not open_("/adminer.php"), "phpMyAdmin/Adminer เข้าถึงได้" if (open_("/phpmyadmin/") or open_("/adminer.php")) else "เข้าถึงไม่ได้")
    path_(147, paths.get("/.well-known/security.txt", {}).get("code") == 200,
         "มี security.txt" if paths.get("/.well-known/security.txt", {}).get("code") == 200 else "ไม่มี security.txt")

    # ── DNS ──────────────────────────────────────────
    if not dns.get("skipped"):
        set_(136, dns.get("spf") and dns.get("dmarc"),
             f"SPF: {'มี' if dns.get('spf') else 'ไม่มี'} / DMARC: {'มี' if dns.get('dmarc') else 'ไม่มี'}")

    # ── Third-party scripts ──────────────────────────
    ed = page.get("external_script_domains", [])
    set_(122, not ed or page.get("scripts_with_sri", 0) > 0,
         f"script ภายนอก {len(ed)} โดเมน, มี integrity {page.get('scripts_with_sri',0)} ตัว")
    r[173] = ("ผ่าน" if ed else "ยังไม่ตรวจ", "โดเมนที่โหลด: " + (", ".join(ed) if ed else "ไม่พบ"))
    if page.get("trackers_on_load"):
        set_(163, False, "pixel/tracker โหลดทันทีตอนเปิดหน้า: " + ", ".join(page["trackers_on_load"]) + " — ต้องยืนยันว่ายิงหลัง consent")

    # ── Local project ────────────────────────────────
    if not proj.get("skipped"):
        npmv = proj.get("npm_vulnerabilities") or {}
        if npmv:
            hi = (npmv.get("high", 0) or 0) + (npmv.get("critical", 0) or 0)
            set_(114, hi == 0, f"npm audit: critical {npmv.get('critical',0)}, high {npmv.get('high',0)}, moderate {npmv.get('moderate',0)}")
        set_(123, bool(proj.get("lockfile_committed")), "พบ lock file" if proj.get("lockfile_committed") else "ไม่พบ lock file")
        set_(133, (proj.get("world_writable_files", 0) or 0) == 0, f"ไฟล์ที่ world-writable: {proj.get('world_writable_files',0)}")
        if proj.get("files_with_possible_secrets") is not None:
            # 102 "API key ไม่ถูกฝังใน frontend / public repo" (Critical), not
            # 113 "CORS ของ API จำกัดเฉพาะโดเมนของบริษัท" (High). This grep
            # reads source files, so it answers 102 — the old mapping put the
            # finding under an unrelated heading AND understated its risk.
            files = proj.get("secret_files") or []
            where = (" — " + ", ".join(files[:5]) + ("…" if len(files) > 5 else "")) if files else ""
            set_(102, proj["files_with_possible_secrets"] == 0,
                 f"ไฟล์ที่มีข้อความคล้าย secret: {proj['files_with_possible_secrets']}{where}"
                 + (" (ตรวจด้วยตาว่าเป็นค่าจริงหรือ fixture ในเทสต์)" if files else ""))

    # Nothing answered on 443.
    #
    # Every verdict above about headers, cookies, redirects and exposed files
    # then rests on silence: "no HSTS header" is true only because there was no
    # response at all, and "no .env exposed" only because nothing was served.
    # A scan that reached nothing has found nothing — say so, rather than hand
    # back a page of findings about a host that was never online.
    if scan.get("meta", {}).get("reachable") is False:
        OFFLINE_OK = {102, 114, 123, 133, 136}  # local project + DNS, still valid
        for i in list(r):
            if i not in OFFLINE_OK:
                unknown_(i, "เว็บไม่ตอบสนองบนพอร์ต 443 จากเครื่องที่รันสแกน — ยังไม่ได้ตรวจอะไรจากตัวเว็บเลย")

    return r


def prepare(scan_path):
    """The scan, scored against the checklist. Shared by every output format —
    two writers with two copies of this would disagree the first time either
    changed."""
    scan = json.load(open(scan_path, encoding="utf-8"))
    items = json.load(open(CHECKLIST, encoding="utf-8"))
    res = evaluate(scan)

    rows = [{**it, **dict(zip(("status", "evidence"), res.get(it["id"], ("ยังไม่ตรวจ", ""))))}
            for it in items]
    fails = sorted([r for r in rows if r["status"] == "ไม่ผ่าน"],
                   key=lambda x: (RISK_ORDER[x["risk"]], x["id"]))

    return {
        "scan": scan,
        "domain": scan.get("meta", {}).get("domain", "-"),
        "rows": rows,
        "fails": fails,
        "auto": [r for r in rows if r["status"] != "ยังไม่ตรวจ"],
    }


def summary(d):
    return {"domain": d["domain"], "auto_checked": len(d["auto"]), "failed": len(d["fails"]),
            "by_risk": {k: sum(1 for f in d["fails"] if f["risk"] == k) for k in RISK_ORDER}}


def build_html(scan_path, out_path):
    """The readable one. Kept in its own module so the spreadsheet writer and
    the page writer cannot drift apart on the scoring — both call prepare()."""
    from _html_report import render

    d = prepare(scan_path)
    with open(out_path, "w", encoding="utf-8") as fh:
        fh.write(render(d))

    return summary(d)


def build(scan_path, out_path):
    d = prepare(scan_path)
    scan, domain, rows, fails, auto = d["scan"], d["domain"], d["rows"], d["fails"], d["auto"]

    wb = Workbook()

    # ── Summary ──────────────────────────────────────
    ws = wb.active; ws.title = "Summary"; ws.sheet_view.showGridLines = False
    for c, w in zip("ABCDE", [3, 34, 16, 16, 40]): ws.column_dimensions[c].width = w
    ws["B2"] = f"รายงานผลตรวจความปลอดภัยเว็บไซต์ — {domain}"
    ws["B2"].font = Font(name=F, size=16, bold=True, color=NAVY)
    ws["B3"] = f"สแกนเมื่อ {scan.get('meta',{}).get('scanned_at','-')}  |  เช็คลิสต์ 189 รายการ  |  ตรวจอัตโนมัติได้ {len(auto)} ข้อ"
    ws["B3"].font = Font(name=F, size=10, color="595959")

    ws["B5"] = "ผลรวมตามระดับความเสี่ยง"
    ws["B5"].font = Font(name=F, size=12, bold=True, color=NAVY)
    for i, t in enumerate(["ระดับ", "ไม่ผ่าน", "ผ่าน", "ยังไม่ตรวจ (ต้องตรวจเอง)"]):
        c = ws.cell(6, 2 + i, t)
        c.font = Font(name=F, size=10, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor="2E5496")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX
    for i, lvl in enumerate(["Critical", "High", "Medium", "Low"]):
        r = 7 + i
        sub = [x for x in rows if x["risk"] == lvl]
        vals = [lvl,
                sum(1 for x in sub if x["status"] == "ไม่ผ่าน"),
                sum(1 for x in sub if x["status"] == "ผ่าน"),
                sum(1 for x in sub if x["status"] == "ยังไม่ตรวจ")]
        for j, v in enumerate(vals):
            c = ws.cell(r, 2 + j, v); c.border = BOX
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.font = Font(name=F, size=10, bold=(j == 0), color=RISK_TEXT[lvl] if j == 0 else "000000")
            if j == 0: c.fill = PatternFill("solid", fgColor=RISK_FILL[lvl])

    ws["B12"] = "3 อันดับที่ควรแก้ก่อน"
    ws["B12"].font = Font(name=F, size=12, bold=True, color=NAVY)
    if fails:
        for i, f in enumerate(fails[:3]):
            r = 13 + i
            ws.cell(r, 2, f"#{f['id']} [{f['risk']}]").font = Font(name=F, size=10, bold=True, color=RISK_TEXT[f["risk"]])
            c = ws.cell(r, 3, f"{f['item']} — {f['evidence']}")
            c.font = Font(name=F, size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
            ws.row_dimensions[r].height = 30
    else:
        ws.cell(13, 2, "ไม่พบข้อที่ไม่ผ่านจากการตรวจอัตโนมัติ").font = Font(name=F, size=10, color="1F6E43")

    ws.cell(18, 2, "หมายเหตุ: 'ยังไม่ตรวจ' ไม่ได้แปลว่าปลอดภัย — เป็นข้อที่สคริปต์ตรวจแทนไม่ได้ ต้องตรวจด้วยคนตามชีต Checklist").font = Font(name=F, size=9, color="808080")

    # ── Findings ─────────────────────────────────────
    fs = wb.create_sheet("Findings")
    heads = ["ลำดับ", "หมวด", "รายการที่ไม่ผ่าน", "ความเสี่ยง", "หลักฐานที่พบ", "ผู้รับผิดชอบ", "กำหนดแก้"]
    for i, (t, w) in enumerate(zip(heads, [7, 28, 58, 13, 52, 18, 14]), start=1):
        c = fs.cell(1, i, t)
        c.font = Font(name=F, size=10, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor="2E5496")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX
        fs.column_dimensions[get_column_letter(i)].width = w
    for j, f in enumerate(fails, start=2):
        for i, v in enumerate([f["id"], f["category"], f["item"], f["risk"], f["evidence"], "", ""], start=1):
            c = fs.cell(j, i, v); c.border = BOX
            c.font = Font(name=F, size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
        fs.cell(j, 4).fill = PatternFill("solid", fgColor=RISK_FILL[f["risk"]])
        fs.cell(j, 4).font = Font(name=F, size=10, bold=True, color=RISK_TEXT[f["risk"]])
        fs.cell(j, 4).alignment = Alignment(horizontal="center", vertical="center")
        fs.row_dimensions[j].height = 30
    if not fails:
        fs.cell(2, 3, "ไม่พบข้อที่ไม่ผ่านจากการตรวจอัตโนมัติ").font = Font(name=F, size=10)
    fs.freeze_panes = "C2"
    fs.auto_filter.ref = f"A1:G{max(2, len(fails)+1)}"

    # ── Checklist ────────────────────────────────────
    cl = wb.create_sheet("Checklist")
    heads = ["ลำดับ", "หมวด", "รายการตรวจสอบ", "ความเสี่ยง", "วิธีตรวจ", "สถานะ", "หลักฐาน / บันทึก"]
    for i, (t, w) in enumerate(zip(heads, [7, 28, 58, 13, 36, 14, 44]), start=1):
        c = cl.cell(1, i, t)
        c.font = Font(name=F, size=10, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor="2E5496")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX
        cl.column_dimensions[get_column_letter(i)].width = w
    for j, x in enumerate(rows, start=2):
        for i, v in enumerate([x["id"], x["category"], x["item"], x["risk"], x["method"], x["status"], x["evidence"]], start=1):
            c = cl.cell(j, i, v); c.border = BOX
            c.font = Font(name=F, size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
        cl.cell(j, 4).fill = PatternFill("solid", fgColor=RISK_FILL[x["risk"]])
        cl.cell(j, 4).font = Font(name=F, size=10, bold=True, color=RISK_TEXT[x["risk"]])
        cl.cell(j, 4).alignment = Alignment(horizontal="center", vertical="center")
        cl.cell(j, 6).fill = PatternFill("solid", fgColor=STAT_FILL.get(x["status"], "FFFFFF"))
        cl.cell(j, 6).alignment = Alignment(horizontal="center", vertical="center")
        cl.row_dimensions[j].height = 28
    cl.freeze_panes = "C2"
    cl.auto_filter.ref = f"A1:G{len(rows)+1}"

    wb.save(out_path)
    return summary(d)


def write(scan_path, out_path):
    """Format follows the extension: .html for reading, .xlsx for handing work
    out. A spreadsheet is a work queue, not a document — the HTML is what you
    actually read the findings in."""
    if out_path.lower().endswith((".html", ".htm")):
        return build_html(scan_path, out_path)
    return build(scan_path, out_path)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("scan_json")
    ap.add_argument("--out", default="Security_Report.html",
                    help="ลงท้าย .html = รายงานอ่านบนเบราว์เซอร์ (ค่าตั้งต้น), .xlsx = ตารางไว้มอบหมายงาน")
    ap.add_argument("--also", default=None, help="ออกอีกฟอร์แมตพร้อมกัน เช่น --also Report.xlsx")
    a = ap.parse_args()
    out = write(a.scan_json, a.out)
    if a.also:
        write(a.scan_json, a.also)
    print(json.dumps(out, ensure_ascii=False, indent=1))
